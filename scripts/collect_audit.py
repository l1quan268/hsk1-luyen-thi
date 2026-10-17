import docx, json, re, sys, os

sys.stdout.reconfigure(encoding='utf-8')

# Load docx tables and paragraphs
dap_an_doc = docx.Document('HSK1_6_De_Dap_An_Transcript_Pinyin.docx')

def normalize_ans(val):
    if not val: return ''
    v = val.strip().upper()
    if v in ('✓', 'T', 'TRUE', 'ĐÚNG'): return 'T'
    if v in ('✗', 'F', 'FALSE', 'SAI'): return 'F'
    return v

# 1. Quick answer tables from Dap An Docx
table_answers = {}
for t_idx, t in enumerate(dap_an_doc.tables):
    eid = t_idx + 1
    table_answers[eid] = {}
    for r in t.rows:
        for c in r.cells:
            for q_num, ans in re.findall(r'(\d+)\s*:\s*([A-F✓✗])', c.text.strip()):
                table_answers[eid][int(q_num)] = normalize_ans(ans)

# 2. Detailed items from Dap An Docx
exam_indices = []
for i, p in enumerate(dap_an_doc.paragraphs):
    m = re.match(r'^ĐỀ\s*0?(\d+)', p.text.strip(), re.I)
    if m: exam_indices.append((int(m.group(1)), i))
exam_indices.append((7, len(dap_an_doc.paragraphs)))

docx_details = {}
for idx in range(len(exam_indices) - 1):
    eid, start_p = exam_indices[idx]
    _, end_p = exam_indices[idx + 1]
    exam_p = [p.text.strip() for p in dap_an_doc.paragraphs[start_p:end_p] if p.text.strip()]
    q_data = {}
    i = 0
    while i < len(exam_p):
        txt = exam_p[i]
        m_q = re.match(r'^0?(\d+)\s*\|\s*Đáp án\s+([A-F✓✗])\s*\|\s*(.*)$', txt, re.DOTALL)
        if m_q:
            qn = int(m_q.group(1))
            ans = normalize_ans(m_q.group(2))
            rest = m_q.group(3).strip().split('\n')
            hz = rest[0].strip() if len(rest) > 0 else ''
            py = rest[1].strip() if len(rest) > 1 else ''
            exp = ''
            dialog_choices = ''
            matched_ans = ''
            fill_word = ''
            i += 1
            while i < len(exam_p):
                nxt = exam_p[i]
                if re.match(r'^0?(\d+)\s*\|\s*Đáp án', nxt) or 'PHẦN ĐỌC' in nxt.upper() or nxt.startswith('ĐỀ'):
                    break
                if nxt.startswith('Giải nghĩa:'): exp = nxt.replace('Giải nghĩa:', '').strip()
                elif nxt.startswith('Lựa chọn:'): dialog_choices = nxt.replace('Lựa chọn:', '').strip()
                elif nxt.startswith('Trả lời:'): matched_ans = nxt.replace('Trả lời:', '').strip()
                elif nxt.startswith('Từ cần điền:'): fill_word = nxt.replace('Từ cần điền:', '').strip()
                elif not py: py = nxt
                i += 1
            q_data[qn] = {
                'answer': ans, 'hanzi': hz, 'pinyin': py, 'explanation': exp,
                'dialog_choices': dialog_choices, 'matched_ans': matched_ans, 'fill_word': fill_word
            }
        else:
            i += 1
    docx_details[eid] = q_data

# 3. Audio scripts
audio_scripts = {}
for eid in range(1, 7):
    fn = f'scripts/audio_transcripts/DE_{eid:02d}_ELEVENLABS_NGAT_TU_NHIEN.txt'
    if os.path.exists(fn):
        with open(fn, 'r', encoding='utf-8') as f:
            lines = [l.strip() for l in f if l.strip()]
        # parse questions from audio script
        # Each question usually starts with "第X题"
        audio_qs = {}
        cur_q = None
        cur_text = []
        for l in lines:
            m = re.match(r'^第([一二三四五六七八九十\d]+)题[。.]?', l)
            if m:
                if cur_q: audio_qs[cur_q] = ' '.join(cur_text)
                # convert chinese num to int
                num_map = {'一':1,'二':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10,
                           '十一':11,'十二':12,'十三':13,'十四':14,'十五':15,'十六':16,'十七':17,'十八':18,'十九':19,'二十':20}
                raw_num = m.group(1)
                cur_q = num_map.get(raw_num, int(raw_num) if raw_num.isdigit() else None)
                cur_text = []
            elif cur_q:
                # ignore pauses like [pause ...]
                cleaned = re.sub(r'\[pause\s+[\d\.]+s\]', '', l).strip()
                if cleaned: cur_text.append(cleaned)
        if cur_q: audio_qs[cur_q] = ' '.join(cur_text)
        audio_scripts[eid] = audio_qs

# 4. Exam JSON and Solutions JSON
exams = {}
solutions = {}
for eid in range(1, 7):
    with open(f'src/data/exam{eid:02d}.json', 'r', encoding='utf-8') as f: exams[eid] = json.load(f)
    with open(f'src/data/solutions{eid:02d}.json', 'r', encoding='utf-8') as f: solutions[eid] = json.load(f)

# Asset semantic descriptions
ASSET_DESC = {
    "img01": "1 con mèo vàng (猫)",
    "img02": "Vé xem phim (看电影 / 电影票 / 电影院)",
    "img03": "Nốt nhạc (听音乐 / 音乐)",
    "img04": "Khuôn mặt nữ tóc nâu (người nữ)",
    "img05": "Tàu điện ngầm trong hầm (地铁)",
    "img06": "1 mặt chó đốm nâu (1 con chó, 狗)",
    "img07": "1 con cá xanh (鱼)",
    "img08": "Tách trà xanh bốc khói (茶 / 喝茶)",
    "img09": "Ly bia bọt (啤酒 / 喝啤酒)",
    "img10": "Tách cà phê nóng (咖啡 / 喝咖啡)",
    "img11": "Trường học có đồng hồ (学校)",
    "img12": "Bệnh viện có chữ thập đỏ (医院)",
    "img13": "Ngân hàng có ký hiệu $ (银行)",
    "img14": "Điện thoại bàn màu đỏ (电话 / 打电话)",
    "img15": "Khung tranh phong cảnh cây cỏ (照片 / bức ảnh)",
    "img16": "Máy tính xách tay laptop (电脑)",
    "img17": "Bát cơm trắng (米饭 / 吃饭)",
    "img18": "Bánh sinh nhật cắm 3 nến KHÔNG có ngày tháng (蛋糕 / 生日)",
    "img19": "Xe buýt (公交车 / 公共汽车)",
    "img20": "Bảng clapperboard đóng phim (看电影 / 电影)",
    "img21": "Chồng 3 quyển sách (3 quyển sách / 书)",
    "img22": "Xe ô tô con màu đỏ (车 / 汽车 / 出租车)",
    "img23": "Ngôi nhà (家 / 房子)",
    "img24": "Xe taxi vàng có sọc caro (出租车)",
    "img25": "Xe đạp (自行车)",
    "img26": "Ghế gỗ (椅子)",
    "img27": "Máy bay đang bay (飞机 / 坐飞机)",
    "img28": "Cây cải thảo / rau xanh (菜 / 白菜)",
    "img29": "Quả táo đỏ có cuống lá (苹果)",
    "img30": "Miếng thịt đỏ có xương (肉)",
    "img31": "Cây xanh công viên (树 / 公园)",
    "img32": "Cửa hàng tiện lợi 24h (商店 / 便利店)",
    "img33": "Đĩa mì spaghetti dĩa gắp (面条 / 吃面条)",
    "img34": "Tô mì ramen có trứng và đũa (面条 / 吃面条)",
    "img35": "Ống nghe y tế bác sĩ (ống nghe, stethoscope)",
    "img36": "Chùm nho tím (葡萄)",
    "img37": "Cặp tài liệu da màu nâu (công văn / đi làm / 爸爸在公司工作)",
    "img38": "Bé trai đang đi bộ (走路 / 走)",
    "img39": "Tờ lịch ghi July 17 KHÔNG ghi thứ (lịch July 17)"
}

print(f"Loaded: 6 exams, {len(docx_details)} docx details, {len(audio_scripts)} audio scripts.")

# Let's inspect all 240 questions and collect audit records
records = []
for eid in range(1, 7):
    ex = exams[eid]
    sol = solutions[eid]
    d_det = docx_details.get(eid, {})
    t_ans_map = table_answers.get(eid, {})
    a_sc = audio_scripts.get(eid, {})
    
    for q in ex['questions']:
        qn = q['questionNumber']
        qtype = q.get('questionType')
        qsec = q.get('section')
        
        # Docx table answer
        doc_tbl_ans = t_ans_map.get(qn, '')
        
        # Docx detail
        doc_item = d_det.get(qn, {})
        doc_dtl_ans = doc_item.get('answer', '')
        
        # Web answer
        web_ans = sol['answers'].get(str(qn), '')
        sol_item = sol['items'].get(str(qn), {})
        
        # Hanzi, Pinyin, Meaning
        hanzi = q.get('hanzi') or sol_item.get('hanzi') or doc_item.get('hanzi') or ''
        pinyin = q.get('pinyin') or sol_item.get('pinyin') or doc_item.get('pinyin') or ''
        exp = sol_item.get('explanation') or doc_item.get('explanation') or ''
        
        # Image IDs
        image_ids = q.get('imageIds', [])
        choice_img_ids = [c.get('imageId') for c in q.get('choices', []) if c.get('imageId')]
        all_imgs = image_ids + choice_img_ids
        
        # Audio text
        audio_txt = a_sc.get(qn, '')
        
        records.append({
            'examId': eid,
            'questionNumber': qn,
            'type': qtype,
            'section': qsec,
            'hanzi': hanzi,
            'pinyin': pinyin,
            'meaning': exp,
            'docTableAns': doc_tbl_ans,
            'docDetailAns': doc_dtl_ans,
            'webAns': web_ans,
            'imageIds': all_imgs,
            'audioScript': audio_txt,
            'choices': q.get('choices', []),
            'solItem': sol_item
        })

print(f"Total question records: {len(records)}")

with open('scripts/raw_audit_records.json', 'w', encoding='utf-8') as f:
    json.dump(records, f, ensure_ascii=False, indent=2)
print("Saved raw_audit_records.json")
