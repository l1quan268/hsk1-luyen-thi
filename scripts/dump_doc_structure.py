import docx
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('HSK1_6_De_Dap_An_Transcript_Pinyin.docx')

print('TABLES COUNT:', len(doc.tables))
for i, t in enumerate(doc.tables):
    print(f'\n=== TABLE {i} ===')
    for r_idx, r in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f'R{r_idx}: ' + ' | '.join(cells))

# Let's inspect paragraphs around ĐỀ 01
print('\n=== FIRST 100 PARAGRAPHS ===')
for i in range(min(100, len(doc.paragraphs))):
    txt = doc.paragraphs[i].text.strip()
    if txt:
        print(f'P{i}: {txt}')
