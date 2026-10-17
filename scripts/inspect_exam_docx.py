import docx, re, json, os, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('HSK1_6_De_Thi_Thu_Full_Pinyin.docx')
rels = {r.rId: r.target_ref for r in doc.part.rels.values() if 'image' in r.reltype}

print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))
print("Total image rels:", len(rels))

for i, p in enumerate(doc.paragraphs):
    if re.match(r'^ĐỀ', p.text.strip(), re.I):
        print(f"P{i}: {p.text.strip()}")
