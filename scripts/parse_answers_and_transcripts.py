import docx
import json
import re
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('HSK1_6_De_Dap_An_Transcript_Pinyin.docx')

def normalize_ans(val):
    v = val.strip().upper()
    if v == '✓':
        return 'T'
    if v == '✗':
        return 'F'
    return v

# 1. Parse tables
table_answers = {}
for t_idx, t in enumerate(doc.tables):
    exam_id = t_idx + 1
    table_answers[exam_id] = {}
    for r in t.rows:
        for c in r.cells:
            txt = c.text.strip()
            matches = re.findall(r'(\d+)\s*:\s*([A-F✓✗])', txt)
            for q_num, ans in matches:
                table_answers[exam_id][int(q_num)] = normalize_ans(ans)

print("Tables summary:")
for eid in range(1, 7):
    print(f"Exam {eid}: {len(table_answers[eid])} answers")

# 2. Parse paragraphs
# Find exam bounds
exam_indices = []
for i, p in enumerate(doc.paragraphs):
    m = re.match(r'^ĐỀ\s*0?(\d+)', p.text.strip(), re.IGNORECASE)
    if m:
        exam_indices.append((int(m.group(1)), i))

exam_indices.append((7, len(doc.paragraphs)))

parsed_exams = {}

for idx in range(len(exam_indices) - 1):
    eid, start_p = exam_indices[idx]
    next_eid, end_p = exam_indices[idx + 1]
    
    exam_p = [p.text.strip() for p in doc.paragraphs[start_p:end_p] if p.text.strip()]
    
    # We will parse questions
    questions_data = {}
    
    # Parse questions from exam_p
    i = 0
    current_section = 'listening'
    while i < len(exam_p):
        txt = exam_p[i]
        if 'PHẦN ĐỌC' in txt.upper():
            current_section = 'reading'
            i += 1
            continue
        if 'PHẦN NGHE' in txt.upper() or 'BẢNG ĐÁP ÁN' in txt.upper() or txt.startswith('ĐỀ'):
            i += 1
            continue
        
        # Check if line starts with question number: e.g. "01 | Đáp án ✗ | ..."
        m_q = re.match(r'^0?(\d+)\s*\|\s*Đáp án\s+([A-F✓✗])\s*\|\s*(.*)$', txt, re.DOTALL)
        if m_q:
            qnum = int(m_q.group(1))
            ans = normalize_ans(m_q.group(2))
            rest = m_q.group(3).strip()
            
            # rest might have hanzi \n pinyin or just hanzi
            lines = [l.strip() for l in rest.split('\n') if l.strip()]
            hanzi = lines[0] if len(lines) > 0 else ''
            pinyin = lines[1] if len(lines) > 1 else ''
            
            q_obj = {
                'questionNumber': qnum,
                'section': 'listening' if qnum <= 20 else 'reading',
                'answer': ans,
                'hanzi': hanzi,
                'pinyin': pinyin,
                'explanation': '',
                'dialogChoices': None,
                'matchedAnswer': None,
                'fillWord': None,
                'completeSentence': None
            }
            
            i += 1
            # Next lines belong to this question until next question or section
            while i < len(exam_p):
                nxt = exam_p[i]
                if re.match(r'^0?(\d+)\s*\|\s*Đáp án', nxt) or 'PHẦN ĐỌC' in nxt.upper() or nxt.startswith('ĐỀ'):
                    break
                
                if nxt.startswith('Giải nghĩa:'):
                    q_obj['explanation'] = nxt.replace('Giải nghĩa:', '').strip()
                elif nxt.startswith('Lựa chọn:'):
                    q_obj['dialogChoices'] = nxt.replace('Lựa chọn:', '').strip()
                elif nxt.startswith('Trả lời:'):
                    # For reading 31-35
                    ans_text = nxt.replace('Trả lời:', '').strip()
                    ans_lines = [l.strip() for l in ans_text.split('\n') if l.strip()]
                    q_obj['matchedAnswer'] = {
                        'hanzi': ans_lines[0] if len(ans_lines) > 0 else '',
                        'pinyin': ans_lines[1] if len(ans_lines) > 1 else ''
                    }
                elif nxt.startswith('Từ cần điền:'):
                    # For reading 36-40
                    fill_text = nxt.replace('Từ cần điền:', '').strip()
                    fill_lines = [l.strip() for l in fill_text.split('\n') if l.strip()]
                    fill_hanzi = fill_lines[0] if len(fill_lines) > 0 else ''
                    fill_pinyin = fill_lines[1] if len(fill_lines) > 1 else ''
                    q_obj['fillWord'] = {
                        'hanzi': fill_hanzi,
                        'pinyin': fill_pinyin
                    }
                else:
                    # In case pinyin was on a separate paragraph
                    if not q_obj['pinyin']:
                        q_obj['pinyin'] = nxt
                i += 1
            
            # For fill bank, compute completeSentence
            if qnum >= 36 and qnum <= 40 and q_obj['fillWord']:
                fh = q_obj['fillWord']['hanzi']
                # replace underscore sequence with filled hanzi
                q_obj['completeSentence'] = re.sub(r'[_—\s]{3,}', f' {fh} ', q_obj['hanzi']).strip()
                q_obj['completeSentence'] = re.sub(r'\s+', ' ', q_obj['completeSentence'])
            
            questions_data[qnum] = q_obj
        else:
            i += 1
            
    parsed_exams[eid] = questions_data

print("\nParsed questions summary:")
for eid in range(1, 7):
    print(f"Exam {eid}: {len(parsed_exams[eid])} questions parsed")
