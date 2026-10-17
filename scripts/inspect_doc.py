import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('HSK1_6_De_Dap_An_Transcript_Pinyin.docx')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')

for i, t in enumerate(doc.tables):
    print(f'\n--- Table {i} (rows={len(t.rows)}, cols={len(t.columns)}) ---')
    for r_idx, r in enumerate(t.rows):
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if r_idx < 10 or r_idx >= len(t.rows) - 2:
            print(f'  Row {r_idx}: ' + ' | '.join(row_txt))

print('\n--- Paragraph Headings / Samples ---')
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('ĐỀ') or txt.startswith('H1') or 'ĐÁP ÁN' in txt.upper() or 'TRANSCRIPT' in txt.upper() or 'PHẦN' in txt.upper():
        print(f'P{i}: {txt}')
