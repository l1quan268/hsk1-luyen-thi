import docx
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('HSK1_6_De_Dap_An_Transcript_Pinyin.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

exam_starts = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if re.match(r'^ĐỀ\s*0?(\d+)', txt, re.IGNORECASE):
        exam_starts.append((i, txt))

print("Exam headings:", exam_starts)

for t_idx, t in enumerate(doc.tables):
    print(f"\nTable {t_idx} has {len(t.rows)} rows")
    # count how many answers in table
    answers = {}
    for r in t.rows:
        for c in r.cells:
            txt = c.text.strip()
            # like '01: ✓' or '01: C'
            matches = re.findall(r'(\d+)\s*:\s*([A-F✓✗])', txt)
            for num, ans in matches:
                answers[int(num)] = ans
    print(f"Table {t_idx} extracted {len(answers)} answers. Sample: {sorted(list(answers.items()))[:5]} ... {sorted(list(answers.items()))[-5:]}")
